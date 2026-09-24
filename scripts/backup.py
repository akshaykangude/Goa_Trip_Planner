#!/usr/bin/env python3
"""
Pulls the shared trip data from the Apps Script backend and writes versioned
CSV + Excel + JSON backups into the repo's backups/ folder.

Run by .github/workflows/backup.yml every 2 hours (and on demand).
Needs two repository secrets:
  TRIP_API_URL   – the Apps Script web-app URL (same as assets/config.js)
  TRIP_PASSCODE  – the trip passcode

Layout written:
  backups/latest/            ← always the newest copy (overwritten)
  backups/history/YYYY-MM-DD/HHMM_revN/   ← one folder per changed version
"""
import csv
import datetime as dt
import io
import json
import os
import pathlib
import sys
import urllib.request

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "backups"
IST = dt.timezone(dt.timedelta(hours=5, minutes=30))


def fetch_export(url: str, key: str) -> dict:
    body = json.dumps({"action": "export", "key": key, "user": "github-backup"}).encode()
    req = urllib.request.Request(url, data=body, headers={"Content-Type": "text/plain;charset=utf-8"})
    # Apps Script answers POST with a 302 to googleusercontent.com; urllib follows it as GET (what Google expects).
    with urllib.request.urlopen(req, timeout=90) as r:
        return json.loads(r.read().decode("utf-8"))


def to_number(v: str):
    try:
        if v.strip() == "":
            return v
        f = float(v.replace(",", ""))
        return int(f) if f.is_integer() and "." not in v else f
    except ValueError:
        return v


def write_xlsx(path: pathlib.Path, tables: dict) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed – skipping .xlsx")
        return False
    wb = Workbook()
    wb.remove(wb.active)
    for name, text in tables.items():
        ws = wb.create_sheet(title=name[:31])
        rows = list(csv.reader(io.StringIO(text.lstrip("﻿"))))
        for i, row in enumerate(rows):
            ws.append(row if i == 0 else [to_number(c) for c in row])
        if rows:
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="FDE7D6")
            ws.freeze_panes = "A2"
            for ci in range(1, len(rows[0]) + 1):
                width = max((len(str(r[ci - 1])) for r in rows if len(r) >= ci), default=8)
                ws.column_dimensions[get_column_letter(ci)].width = min(max(10, width + 2), 60)
    wb.save(path)
    return True


def write_set(folder: pathlib.Path, data: dict) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    for name, text in (data.get("csv") or {}).items():
        safe = "".join(ch for ch in name if ch.isalnum() or ch in " -_").strip().replace(" ", "-")
        (folder / f"{safe}.csv").write_text("﻿" + text, encoding="utf-8")
    (folder / "trip.json").write_text(json.dumps(data.get("state"), ensure_ascii=False, indent=1), encoding="utf-8")
    write_xlsx(folder / "Goa-Trip.xlsx", data.get("csv") or {})


def main() -> int:
    url, key = os.environ.get("TRIP_API_URL", "").strip(), os.environ.get("TRIP_PASSCODE", "").strip()
    if not url or not key:
        print("ℹ️  TRIP_API_URL / TRIP_PASSCODE secrets not set – nothing to back up yet (see docs/SETUP.md step 4).")
        return 0
    data = fetch_export(url, key)
    if not data.get("ok"):
        print("❌ Backend error:", data.get("error"))
        return 1
    rev = int(data.get("rev") or 0)
    if not rev or not data.get("state"):
        print("ℹ️  Backend has no trip data yet.")
        return 0
    latest = OUT / "latest"
    rev_file = latest / "rev.txt"
    if rev_file.exists() and rev_file.read_text().strip() == str(rev):
        print(f"✅ No changes since rev {rev} – skipping.")
        return 0
    now = dt.datetime.now(IST)
    hist = OUT / "history" / now.strftime("%Y-%m-%d") / f"{now.strftime('%H%M')}_rev{rev}"
    write_set(hist, data)
    write_set(latest, data)
    rev_file.write_text(str(rev))
    state = data["state"]
    total = sum(float(e.get("amt") or 0) for e in state.get("expenses", []))
    (latest / "INFO.md").write_text(
        f"# Latest backup\n\n- Revision: **{rev}**\n- Backed up: {now.strftime('%d %b %Y, %I:%M %p IST')}\n"
        f"- Last edit by: {data.get('updatedBy') or '—'} at {data.get('updatedAt') or '—'}\n"
        f"- Bills: {len(state.get('expenses', []))} · Total spent: ₹{total:,.0f}\n\n"
        "Open `Goa-Trip.xlsx` in Excel, or any `.csv` in Excel / Google Sheets.\n"
        "To restore: website → 🕘 History → pick a version, or import `trip.json` (⬆ Import).\n",
        encoding="utf-8")
    print(f"✅ Backed up rev {rev} → {hist.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
